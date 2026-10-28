"""
Document Parser for PrivateDesk / IST Bot.
Extracts clean, structured text and Markdown from PDFs and PowerPoint presentations.
Optimized for local LLM inference (Qwen 2.5):
- Converts PDF and PPTX tables into clean Markdown tables.
- Preserves logical reading order, headings, and bullet points.
- Automatically repairs hyphenated line wraps and fragmented sentences.
- Primary engine: PyMuPDF (fast, layout & table aware).
- Resilient fallback: pypdf (pure Python layout extraction).
"""

from __future__ import annotations

import os
import re
from html.parser import HTMLParser
from urllib.parse import urlparse, unquote
import requests

# Primary engine: PyMuPDF (fast, detects tables and layout reading order)
try:
    import pymupdf
    if hasattr(pymupdf, "no_recommend_layout"):
        pymupdf.no_recommend_layout()
    HAS_PYMUPDF = True
except ImportError:
    pymupdf = None
    HAS_PYMUPDF = False

# Resilient fallback engine: pypdf
from pypdf import PdfReader
from pypdf.errors import PdfReadError, FileNotDecryptedError


def clean_markdown_for_llm(text: str) -> str:
    """
    Cleans raw extracted text into clean, structured Markdown suitable for Qwen:
    - Normalizes control characters and multiple newlines.
    - Rejoins hyphenated line breaks (e.g., 'connec-\ntion' -> 'connection').
    - Standardizes bullet characters (•, ·, etc.) into '- '.
    - Rejoins broken sentences while preserving Markdown tables, headers, and lists.
    """
    if not text:
        return ""

    # Normalize newlines and strip non-printable ASCII control chars
    text = text.replace("\x0c", "\n").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\x00-\x08\x0b\x0e-\x1f\x7f]", "", text)

    # Rejoin hyphenated line breaks
    text = re.sub(r"(\w+)-\n(\w+)", r"\1\2", text)

    # Standardize bullet points
    text = re.sub(r"^[ \t]*[•·▪▸‣⁃]\s*", "- ", text, flags=re.MULTILINE)

    # Process line-by-line to join wrapped sentences within paragraphs
    lines = text.split("\n")
    formatted = []
    for line in lines:
        stripped = line.strip()
        if not stripped:
            if formatted and formatted[-1] != "":
                formatted.append("")
            continue

        is_structural = (
            stripped.startswith(("|", "#", "-", "*", ">")) or
            bool(re.match(r"^\d+[\.\)]\s", stripped))
        )

        if formatted and formatted[-1] != "":
            prev = formatted[-1]
            prev_structural = (
                prev.startswith(("|", "#", "-", "*", ">")) or
                bool(re.match(r"^\d+[\.\)]\s", prev))
            )
            # Rejoin non-structural sentences that wrap across lines
            if not is_structural and not prev_structural and not prev.endswith((".", "!", "?", ":", ";")):
                formatted[-1] = prev + " " + stripped
                continue

        formatted.append(stripped)

    result = "\n".join(formatted).strip()
    return re.sub(r"\n{3,}", "\n\n", result)


class DocumentParser:
    @staticmethod
    def _parse_pdf_with_pymupdf(file_path: str, filename: str) -> list[dict]:
        """Extracts text and markdown tables from PDF using PyMuPDF."""
        doc = pymupdf.open(file_path)
        try:
            if doc.is_encrypted:
                auth = doc.authenticate("")
                if not auth:
                    raise ValueError(f"'{filename}' is password-protected. Please remove the password and retry.")

            total_pages = len(doc)
            total_text_length = 0
            pages_content = []

            for idx, page in enumerate(doc, start=1):
                page_items = []
                used_tables = set()

                # 1. Detect tables on the page
                try:
                    tabs = page.find_tables()
                    tab_list = tabs.tables if tabs else []
                except Exception:
                    tab_list = []

                def block_overlaps_table(b_bbox, t_bbox):
                    bx0, by0, bx1, by1 = b_bbox
                    tx0, ty0, tx1, ty1 = t_bbox
                    cx, cy = (bx0 + bx1) / 2, (by0 + by1) / 2
                    return tx0 <= cx <= tx1 and ty0 <= cy <= ty1

                # 2. Extract text blocks in reading order
                try:
                    blocks = page.get_text("blocks")
                except Exception:
                    blocks = []

                for b in blocks:
                    bbox = b[:4]
                    b_text = b[4].strip()
                    if not b_text:
                        continue

                    # Check if block is inside a table
                    in_table = False
                    for t_idx, t in enumerate(tab_list):
                        if block_overlaps_table(bbox, t.bbox):
                            in_table = True
                            if t_idx not in used_tables:
                                used_tables.add(t_idx)
                                md_table = t.to_markdown().strip()
                                if md_table:
                                    page_items.append((t.bbox[1], md_table))
                            break

                    if not in_table:
                        page_items.append((bbox[1], b_text))

                # Insert any tables that did not intersect with text blocks
                for t_idx, t in enumerate(tab_list):
                    if t_idx not in used_tables:
                        md_table = t.to_markdown().strip()
                        if md_table:
                            page_items.append((t.bbox[1], md_table))

                # Sort all content vertically (Y-coordinate reading order)
                page_items.sort(key=lambda x: x[0])
                raw_page_text = "\n\n".join(content for _, content in page_items)

                # Clean into structured markdown for Qwen
                cleaned = clean_markdown_for_llm(raw_page_text)
                if cleaned:
                    total_text_length += len(cleaned)
                    pages_content.append({
                        "filename": filename,
                        "page": idx,
                        "text": cleaned
                    })

            if total_pages > 0 and total_text_length < 50:
                raise ValueError(
                    f"'{filename}' appears to be a scanned document or image without selectable text. "
                    "Please upload a digital or OCR-enabled PDF."
                )

            return pages_content
        finally:
            doc.close()

    @staticmethod
    def _parse_pdf_with_pypdf(file_path: str, filename: str) -> list[dict]:
        """Resilient pure-Python fallback using pypdf."""
        reader = PdfReader(file_path)
        if reader.is_encrypted:
            try:
                decrypted = reader.decrypt("")
                if not decrypted:
                    raise ValueError(f"'{filename}' is password-protected. Please remove the password and retry.")
            except Exception:
                raise ValueError(f"'{filename}' is password-protected. Please remove the password and retry.")

        total_pages = len(reader.pages)
        total_text_length = 0
        pages_content = []

        for idx, page in enumerate(reader.pages, start=1):
            raw_text = ""
            try:
                raw_text = page.extract_text(extraction_mode="layout") or ""
            except Exception:
                try:
                    raw_text = page.extract_text() or ""
                except Exception:
                    raw_text = ""

            cleaned = clean_markdown_for_llm(raw_text)
            if cleaned:
                total_text_length += len(cleaned)
                pages_content.append({
                    "filename": filename,
                    "page": idx,
                    "text": cleaned
                })

        if total_pages > 0 and total_text_length < 50:
            raise ValueError(
                f"'{filename}' appears to be a scanned document or image without selectable text. "
                "Please upload a digital or OCR-enabled PDF."
            )

        return pages_content

    @classmethod
    def parse_pdf(cls, file_path: str, filename: str) -> list[dict]:
        """
        Extracts structured text and Markdown from a PDF file page-by-page.
        Uses PyMuPDF if available for table detection and layout preservation,
        with seamless fallback to pypdf layout mode.
        """
        if HAS_PYMUPDF:
            try:
                return cls._parse_pdf_with_pymupdf(file_path, filename)
            except ValueError:
                # Password errors or scanned document errors should be raised directly
                raise
            except Exception:
                # If pymupdf encounters unexpected low-level corruption, try pypdf fallback
                pass

        try:
            return cls._parse_pdf_with_pypdf(file_path, filename)
        except (FileNotDecryptedError, ValueError):
            raise
        except PdfReadError as e:
            raise ValueError(f"Failed to read PDF '{filename}': {str(e)}")
        except Exception as e:
            raise ValueError(f"Error parsing '{filename}': {str(e)}")

    @staticmethod
    def parse_pptx(file_path: str, filename: str) -> list[dict]:
        """
        Extracts text from a PowerPoint (.pptx) presentation slide by slide.
        Extracts slide headings, bullet points, shape text, PowerPoint tables (as Markdown),
        and speaker notes.
        Returns a list of dicts:
        [{"filename": filename, "page": slide_number, "text": "Slide text..."}]
        """
        try:
            from pptx import Presentation
        except ImportError:
            raise ValueError("python-pptx is required to parse PowerPoint files. Run: pip install python-pptx")

        slides_content = []
        try:
            prs = Presentation(file_path)
            def _iter_shapes(shapes):
                for shp in shapes:
                    if hasattr(shp, "shapes"):
                        yield from _iter_shapes(shp.shapes)
                    else:
                        yield shp

            for slide_num, slide in enumerate(prs.slides, start=1):
                slide_text = []

                # Extract text and tables from slide shapes (including grouped shapes)
                for shape in _iter_shapes(slide.shapes):
                    if getattr(shape, "has_table", False):
                        rows_data = []
                        for row in shape.table.rows:
                            rows_data.append([cell.text.strip().replace("\n", " ") for cell in row.cells])
                        if rows_data:
                            header = "| " + " | ".join(rows_data[0]) + " |"
                            sep = "| " + " | ".join(["---"] * len(rows_data[0])) + " |"
                            body_rows = ["| " + " | ".join(r) + " |" for r in rows_data[1:]]
                            md_table = "\n".join([header, sep] + body_rows)
                            slide_text.append(md_table)
                    elif hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text.strip())

                # Extract speaker notes if available
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        slide_text.append(f"[Speaker Notes: {notes}]")

                cleaned_text = clean_markdown_for_llm("\n\n".join(slide_text))
                if cleaned_text:
                    slides_content.append({
                        "filename": filename,
                        "page": slide_num,
                        "text": cleaned_text
                    })

            if not slides_content:
                raise ValueError(f"'{filename}' contains no readable text across its slides.")

            return slides_content

        except Exception as e:
            if isinstance(e, ValueError):
                raise
            raise ValueError(f"Error reading PowerPoint '{filename}': {str(e)}")

    @staticmethod
    def parse_url(url: str, html_content: str | None = None) -> dict:
        """
        Fetches and extracts clean Markdown sections from a web page or Confluence article.
        If a URL has an anchor hash (#topic), selectively extracts ONLY that section.
        """
        parsed_url = urlparse(url)
        fragment = parsed_url.fragment.strip().lower() if parsed_url.fragment else None
        base_url = f"{parsed_url.scheme}://{parsed_url.netloc}{parsed_url.path}"
        if parsed_url.query:
            base_url = f"{base_url}?{parsed_url.query}"

        if html_content is None:
            try:
                headers = {
                    "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36 PrivateDesk/2.0"
                }
                res = requests.get(url, headers=headers, timeout=15, stream=True)
                try:
                    if res.status_code in (401, 403):
                        raise ValueError(
                            f"Access Denied (HTTP {res.status_code}) for '{url}'. "
                            "If this is a private Confluence wiki or intranet page, it must be accessible from your current network/VPN without browser session login."
                        )
                    elif res.status_code == 404:
                        raise ValueError(f"Page not found (HTTP 404): '{url}'.")
                    elif res.status_code != 200:
                        raise ValueError(f"HTTP error {res.status_code} when fetching '{url}'.")

                    content_type = res.headers.get("Content-Type", "").lower()
                    if content_type and not any(t in content_type for t in ("html", "xml", "text/plain")):
                        hint = " Download it and use /upload <file path> instead." if ("pdf" in content_type or "presentation" in content_type) else ""
                        raise ValueError(f"'{url}' is not a web page (content type: {content_type.split(';')[0]}).{hint}")

                    max_bytes = 5 * 1024 * 1024  # 5 MB is far beyond any real article
                    body = bytearray()
                    for piece in res.iter_content(chunk_size=65536):
                        body.extend(piece)
                        if len(body) > max_bytes:
                            raise ValueError(f"'{url}' is too large to index as a web page (over 5 MB).")

                    encoding = res.encoding if res.encoding and "charset" in content_type else None
                    html_text = bytes(body).decode(encoding or "utf-8", errors="replace")
                finally:
                    res.close()
            except requests.exceptions.Timeout:
                raise ValueError(f"Connection timed out while trying to reach '{url}'.")
            except requests.exceptions.ConnectionError:
                raise ValueError(f"Could not connect to '{url}'. Check your internet connection or URL.")
            except Exception as e:
                if isinstance(e, ValueError):
                    raise
                raise ValueError(f"Error fetching '{url}': {str(e)}")
        else:
            html_text = html_content

        html_parser = ConfluenceHTMLParser(target_anchor=fragment)
        html_parser.feed(html_text)
        html_parser.close()

        # Clean document title
        raw_title = html_parser.doc_title.strip()
        if raw_title:
            clean_title = re.sub(
                r"\s*[-|–—]\s*(Confluence|Notion|Wiki|Documentation|Docs|Knowledge Base).*$",
                "",
                raw_title,
                flags=re.IGNORECASE
            ).strip()
            doc_title = clean_title or raw_title
        else:
            path_parts = [p for p in parsed_url.path.split("/") if p]
            doc_title = unquote(path_parts[-1]) if path_parts else parsed_url.netloc

        # Clean section contents
        extracted_sections = []
        for sec in html_parser.sections:
            sec_title = sec["title"].strip()
            sec_text = clean_markdown_for_llm("\n\n".join(sec["lines"]))
            if sec_text:
                extracted_sections.append({
                    "title": sec_title,
                    "anchors": sec["anchors"],
                    "text": sec_text
                })

        if not extracted_sections:
            raise ValueError(f"'{url}' contains no readable article text or paragraphs.")

        # Targeted Anchor Matching
        is_targeted = False
        target_not_found = False
        final_sections = extracted_sections

        if fragment:
            clean_target = fragment.lower().lstrip("#")
            slug_target = re.sub(r"[^a-z0-9]+", "-", clean_target).strip("-")

            matched = []
            for s in extracted_sections:
                s_anchors = [a.lower() for a in s["anchors"]]
                s_title_slug = re.sub(r"[^a-z0-9]+", "-", s["title"].lower()).strip("-")

                # 1. Exact anchor id match
                if clean_target in s_anchors or any(clean_target in a for a in s_anchors):
                    matched.append(s)
                    break
                # 2. Slugified heading match
                elif slug_target and (slug_target == s_title_slug or slug_target in s_title_slug):
                    matched.append(s)
                    break
                # 3. Substring match on title
                elif clean_target in s["title"].lower():
                    matched.append(s)
                    break

            if matched:
                is_targeted = True
                final_sections = matched
            else:
                target_not_found = True

        return {
            "doc_title": doc_title,
            "url": url,
            "base_url": base_url,
            "target_anchor": fragment,
            "is_targeted": is_targeted,
            "target_not_found": target_not_found,
            "sections": [
                {
                    "section_num": idx,
                    "title": s["title"],
                    "anchor": s["anchors"][0] if s["anchors"] else None,
                    "text": s["text"],
                    "url": f"{base_url}#{s['anchors'][0]}" if s["anchors"] else base_url
                }
                for idx, s in enumerate(final_sections, 1)
            ]
        }


class ConfluenceHTMLParser(HTMLParser):
    def __init__(self, target_anchor: str | None = None):
        super().__init__()
        self.target_anchor = target_anchor.lower().lstrip("#") if target_anchor else None

        self.doc_title = ""
        self.in_title = False

        # Tags to ignore completely (tracked via stack so nested divs/tags never desync)
        self.ignored_tags = {
            "script", "style", "nav", "header", "footer", "aside",
            "noscript", "svg", "form", "button", "iframe"
        }
        self.void_tags = {
            "area", "base", "br", "col", "embed", "hr", "img",
            "input", "link", "meta", "param", "source", "track", "wbr"
        }
        self.block_tags = {
            "p", "div", "section", "article", "main", "blockquote",
            "pre", "dl", "dt", "dd", "hr", "br"
        }
        self.ignore_stack: list[str] = []

        # Heading tracking
        self.heading_tags = {"h1", "h2", "h3", "h4"}
        self.current_heading_tag = None
        self.current_heading_text = []
        self.current_heading_anchors = []

        # Sections list: [ {"title": ..., "anchors": [...], "lines": [...]} ]
        self.sections = []
        self.current_section = {
            "title": "Overview",
            "anchors": [],
            "lines": []
        }
        self.current_block: list[str] = []

        # Table tracking
        self.in_table = False
        self.table_rows = []
        self.current_row = []
        self.current_cell = []

        # List tracking
        self.list_depth = 0
        self.in_list_item = False
        self.current_li = []

    def _flush_block(self):
        """Flushes accumulated inline paragraph/block text as a single line."""
        if self.current_block:
            block_text = re.sub(r"\s+", " ", "".join(self.current_block)).strip()
            if block_text:
                self.current_section["lines"].append(block_text)
            self.current_block = []

    def handle_starttag(self, tag, attrs):
        if self.ignore_stack:
            if tag not in self.void_tags:
                self.ignore_stack.append(tag)
            return

        attrs_dict = dict(attrs)
        tag_id = attrs_dict.get("id") or attrs_dict.get("name")
        classes = (attrs_dict.get("class") or "").lower()

        # Ignore Confluence/wiki sidebars, TOC macros, and breadcrumb navigation
        if tag in self.ignored_tags or any(
            c in classes for c in ["toc-macro", "table-of-contents", "breadcrumbs", "page-metadata", "header-actions"]
        ):
            if tag not in self.void_tags:
                self.ignore_stack.append(tag)
            return

        if tag == "title":
            self.in_title = True
            return

        if tag in self.heading_tags:
            self._flush_block()
            self.current_heading_tag = tag
            self.current_heading_text = []
            self.current_heading_anchors = []
            if tag_id:
                self.current_heading_anchors.append(tag_id.lower())
            return

        # Check for anchor identifiers in non-heading tags
        if tag_id:
            if self.current_heading_tag:
                self.current_heading_anchors.append(tag_id.lower())
            else:
                self.current_section["anchors"].append(tag_id.lower())

        if tag == "table":
            self._flush_block()
            self.in_table = True
            self.table_rows = []

        elif tag == "tr" and self.in_table:
            self.current_row = []

        elif tag in ("td", "th") and self.in_table:
            self.current_cell = []

        elif tag in ("ul", "ol"):
            self._flush_block()
            if self.in_list_item and self.current_li:
                li_text = re.sub(r"\s+", " ", "".join(self.current_li)).strip()
                if li_text:
                    indent = "  " * max(0, self.list_depth - 1)
                    self.current_section["lines"].append(f"{indent}- {li_text}")
                self.current_li = []
            self.list_depth += 1

        elif tag == "li":
            self._flush_block()
            self.in_list_item = True
            self.current_li = []

        elif tag in self.block_tags and not self.in_table and not self.in_list_item:
            self._flush_block()

    def handle_endtag(self, tag):
        if self.ignore_stack:
            if tag in self.ignore_stack:
                idx = len(self.ignore_stack) - 1 - self.ignore_stack[::-1].index(tag)
                del self.ignore_stack[idx:]
            return

        if tag == "title":
            self.in_title = False
            return

        if tag in self.heading_tags and tag == self.current_heading_tag:
            heading_title = re.sub(r"\s+", " ", "".join(self.current_heading_text)).strip()
            if heading_title:
                self._flush_block()
                if self.current_section["lines"] or self.current_section["title"] != "Overview":
                    self.sections.append(self.current_section)

                self.current_section = {
                    "title": heading_title,
                    "anchors": list(self.current_heading_anchors),
                    "lines": []
                }
            self.current_heading_tag = None
            self.current_heading_text = []
            self.current_heading_anchors = []

        elif tag == "table" and self.in_table:
            self.in_table = False
            if self.table_rows:
                header = self.table_rows[0]
                md_header = "| " + " | ".join(header) + " |"
                md_sep = "| " + " | ".join(["---"] * len(header)) + " |"
                md_body = ["| " + " | ".join(r) + " |" for r in self.table_rows[1:]]
                table_md = "\n".join([md_header, md_sep] + md_body)
                self.current_section["lines"].append(table_md)
            self.table_rows = []

        elif tag == "tr" and self.in_table:
            if self.current_row:
                self.table_rows.append(self.current_row)
            self.current_row = []

        elif tag in ("td", "th") and self.in_table:
            cell_text = re.sub(r"\s+", " ", "".join(self.current_cell)).strip()
            self.current_row.append(cell_text)
            self.current_cell = []

        elif tag in ("ul", "ol"):
            if self.list_depth > 0:
                self.list_depth -= 1

        elif tag == "li" and self.in_list_item:
            self.in_list_item = False
            li_text = re.sub(r"\s+", " ", "".join(self.current_li)).strip()
            if li_text:
                indent = "  " * max(0, self.list_depth - 1)
                self.current_section["lines"].append(f"{indent}- {li_text}")
            self.current_li = []

        elif tag in self.block_tags and not self.in_table and not self.in_list_item:
            self._flush_block()
            if tag == "p":
                self.current_section["lines"].append("")

    def handle_data(self, data):
        if self.ignore_stack:
            return

        if self.in_title:
            self.doc_title += data
            return

        if self.current_heading_tag:
            self.current_heading_text.append(data)
            return

        if self.in_table and self.current_cell is not None:
            self.current_cell.append(data)
            return

        if self.in_list_item:
            self.current_li.append(data)
            return

        if data.strip():
            self.current_block.append(data)
        elif self.current_block:
            self.current_block.append(" ")

    def close(self):
        super().close()
        self._flush_block()
        if self.current_section["lines"] or self.current_section["title"] != "Overview":
            self.sections.append(self.current_section)


def parse_uploaded_file(file_path_or_url: str, filename: str | None = None) -> list[dict]:
    """Factory function to parse uploaded documents based on extension (.pdf or .pptx) or web URL."""
    if file_path_or_url.startswith(("http://", "https://")):
        web_data = DocumentParser.parse_url(file_path_or_url)
        return [
            {
                "filename": f"[Web] {web_data['doc_title']}" if not web_data["is_targeted"] else f"[Web] {web_data['doc_title']} - {s['title']}",
                "page": s["section_num"],
                "section_title": s["title"],
                "url": s["url"],
                "text": f"## {s['title']}\n\n{s['text']}"
            }
            for s in web_data["sections"]
        ]

    fname = filename or os.path.basename(file_path_or_url)
    ext = os.path.splitext(fname)[1].lower()
    if ext == ".pdf":
        return DocumentParser.parse_pdf(file_path_or_url, fname)
    elif ext == ".pptx":
        return DocumentParser.parse_pptx(file_path_or_url, fname)
    else:
        raise ValueError(f"Unsupported file format '{ext}'. Supported formats: .pdf, .pptx, or web URLs (http/https)")

